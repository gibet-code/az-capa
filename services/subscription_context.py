"""Business Context models, validation, and manual-file resolution."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import PurePath
from typing import Any, Iterable

from core import settings
from core.identity import UserAuditMetadata, for_backend_operation
from clients.resource_graph import get_subscription_inventory
from services.app_scope import resolve_app_scope
from storage.subscription_context_store import (
    ContextFieldConflictError,
    SubscriptionContextStore,
    store as context_store,
)

_FIELD_KEY_RE = re.compile(r"^[a-z][a-z0-9_]{1,31}$")
_RESERVED_FIELD_KEYS = {
    "location",
    "resource_group",
    "subscription_id",
    "subscription_name",
    "unknown",
}
_ZERO_SUBSCRIPTION_ID = "00000000-0000-0000-0000-000000000000"
_REQUIRED_COLUMNS = ("SubscriptionId", "Value")
_VALUE_SOURCE_MANUAL = "manual_file"
_VALUE_SOURCE_MANAGEMENT_GROUP = "management_group_level"
_VALUE_SOURCE_SUBSCRIPTION_TAG = "subscription_tag"
_MG_SELECTED_LEVEL = "selected_level"
_MG_HIERARCHY_PATH = "hierarchy_path"


@dataclass(frozen=True)
class ContextField:
    key: str
    name: str
    enabled: bool
    value_source: dict[str, Any]
    revision: int | None = None


def parse_context_field(payload: Any) -> ContextField:
    """Validate and normalize a context-field definition."""
    if not isinstance(payload, dict):
        raise ValueError("Context field must be a JSON object.")

    key = str(payload.get("key") or "").strip().lower()
    if not _FIELD_KEY_RE.fullmatch(key):
        raise ValueError(
            "Field key must start with a lowercase letter, contain only lowercase "
            "letters, digits, or underscores, and be 2 to 32 characters long."
        )
    if key in _RESERVED_FIELD_KEYS:
        raise ValueError(f"Field key {key!r} is reserved by the application.")

    name = str(payload.get("name") or "").strip()
    if not 1 <= len(name) <= 80:
        raise ValueError("Name must contain between 1 and 80 characters.")

    enabled = payload.get("enabled")
    if not isinstance(enabled, bool):
        raise ValueError("Enabled must be a boolean.")

    value_source = payload.get("valueSource")
    if not isinstance(value_source, dict):
        raise ValueError("Value source must be an object.")
    source_type = value_source.get("type")
    if source_type not in {
        _VALUE_SOURCE_MANUAL,
        _VALUE_SOURCE_MANAGEMENT_GROUP,
        _VALUE_SOURCE_SUBSCRIPTION_TAG,
    }:
        raise ValueError(f"Unsupported Value source type: {source_type!r}.")
    if value_source.get("version") != 1:
        raise ValueError("Value source version must be 1.")
    if source_type == _VALUE_SOURCE_MANAGEMENT_GROUP:
        config = value_source.get("config")
        if not isinstance(config, dict):
            raise ValueError("Management-group Value source config must be an object.")
        level = config.get("level")
        if not isinstance(level, int) or isinstance(level, bool) or not 1 <= level <= 6:
            raise ValueError("Management-group level must be an integer from 1 through 6.")
        if config.get("outputMode") not in {_MG_SELECTED_LEVEL, _MG_HIERARCHY_PATH}:
            raise ValueError("Management-group outputMode must be selected_level or hierarchy_path.")
    if source_type == _VALUE_SOURCE_SUBSCRIPTION_TAG:
        config = value_source.get("config")
        if not isinstance(config, dict):
            raise ValueError("Subscription-tag Value source config must be an object.")
        tag_key = config.get("tagKey")
        if not isinstance(tag_key, str) or not 1 <= len(tag_key.strip()) <= 512:
            raise ValueError("Subscription tagKey must contain between 1 and 512 characters.")

    revision = payload.get("revision")
    if revision is not None and (not isinstance(revision, int) or revision < 1):
        raise ValueError("Revision must be a positive integer when provided.")

    return ContextField(key, name, enabled, value_source, revision)


def parse_manual_mapping(content: bytes, filename: str) -> dict[str, str]:
    """Parse a UTF-8 CSV or XLSX mapping into normalized subscription keys."""
    extension = PurePath(filename or "").suffix.lower()
    if extension == ".csv":
        rows = _csv_rows(content)
    elif extension == ".xlsx":
        rows = _xlsx_rows(content)
    else:
        raise ValueError("Mapping file must use the .csv or .xlsx extension.")
    return _mapping_from_rows(rows)


def _csv_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV mapping file must be UTF-8 encoded.") from exc
    try:
        return list(csv.DictReader(io.StringIO(text, newline="")))
    except csv.Error as exc:
        raise ValueError(f"Invalid CSV mapping file: {exc}") from exc


def _xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - deployment dependency guard
        raise RuntimeError("XLSX support requires the openpyxl package.") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=False)
        header_cells = next(values, ())
        headers = [str(cell.value or "").strip() for cell in header_cells]
        rows: list[dict[str, Any]] = []
        for cells in values:
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("XLSX mapping file must not contain formulas.")
            rows.append({header: cell.value for header, cell in zip(headers, cells)})
        return rows
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - parser exceptions differ by workbook defect
        raise ValueError(f"Invalid XLSX mapping file: {exc}") from exc


def _mapping_from_rows(rows: list[dict[str, Any]]) -> dict[str, str]:
    if not rows:
        raise ValueError("Mapping file contains no data rows.")

    available = {str(column).strip().lower(): str(column) for column in rows[0]}
    missing = [column for column in _REQUIRED_COLUMNS if column.lower() not in available]
    if missing:
        raise ValueError(f"Mapping file is missing required column(s): {missing}.")

    subscription_column = available["subscriptionid"]
    value_column = available["value"]
    mapping: dict[str, str] = {}
    seen: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        subscription_id = str(row.get(subscription_column) or "").strip().lower()
        if not subscription_id:
            continue
        if not _is_subscription_id(subscription_id) or subscription_id == _ZERO_SUBSCRIPTION_ID:
            raise ValueError(f"Invalid SubscriptionId on row {row_number}: {subscription_id!r}.")
        if subscription_id in seen:
            raise ValueError(f"Duplicate SubscriptionId on row {row_number}: {subscription_id}.")
        seen.add(subscription_id)

        value = str(row.get(value_column) or "").strip()
        if len(value) > 256:
            raise ValueError(f"Value on row {row_number} exceeds 256 characters.")
        if value:
            mapping[subscription_id] = value
    return mapping


def _is_subscription_id(value: str) -> bool:
    return bool(re.fullmatch(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", value))


def build_preview(
    mapping: dict[str, str],
    accessible_subscription_ids: Iterable[str],
    in_scope_subscription_ids: Iterable[str],
    missing_reasons: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Return raw values plus in-scope and outside-scope mapping summaries."""
    accessible = _normalized_ids(accessible_subscription_ids)
    in_scope = accessible & _normalized_ids(in_scope_subscription_ids)
    outside_scope = accessible - in_scope
    return {
        "values": {subscription_id: mapping.get(subscription_id) for subscription_id in sorted(accessible)},
        "inScope": _population_summary(mapping, in_scope, missing_reasons or {}),
        "outsideScope": _population_summary(mapping, outside_scope, missing_reasons or {}) if outside_scope else None,
    }


def _normalized_ids(values: Iterable[str]) -> set[str]:
    return {str(value).strip().lower() for value in values if str(value).strip()}


def _population_summary(
    mapping: dict[str, str],
    subscription_ids: set[str],
    missing_reasons: dict[str, str],
) -> dict[str, Any]:
    counts: dict[str, int] = {}
    mapped = 0
    for subscription_id in sorted(subscription_ids):
        value = mapping.get(subscription_id)
        if value is None:
            continue
        mapped += 1
        counts[value] = counts.get(value, 0) + 1
    total = len(subscription_ids)
    unmapped_by_reason: dict[str, int] = {}
    for subscription_id in subscription_ids:
        if subscription_id in mapping:
            continue
        reason = missing_reasons.get(subscription_id, "no_source_value")
        unmapped_by_reason[reason] = unmapped_by_reason.get(reason, 0) + 1
    return {
        "totalSubscriptions": total,
        "mappedSubscriptions": mapped,
        "unmappedSubscriptions": total - mapped,
        "coveragePercent": round(mapped * 100 / total, 1) if total else 0.0,
        "valueCounts": [
            {"value": value, "subscriptionCount": count}
            for value, count in counts.items()
        ],
        "unmappedByReason": unmapped_by_reason,
    }


def build_management_group_mapping(
    inventory: Iterable[dict[str, Any]],
    level: int,
    output_mode: str,
) -> tuple[dict[str, str], dict[str, str]]:
    """Resolve subscription values from management-group ancestor chains."""
    if not isinstance(level, int) or isinstance(level, bool) or not 1 <= level <= 6:
        raise ValueError("Management-group level must be an integer from 1 through 6.")
    if output_mode not in {_MG_SELECTED_LEVEL, _MG_HIERARCHY_PATH}:
        raise ValueError("Management-group outputMode must be selected_level or hierarchy_path.")

    mapping: dict[str, str] = {}
    missing_reasons: dict[str, str] = {}
    no_path: list[str] = []
    for subscription in inventory:
        subscription_id = str(subscription.get("id") or "").strip().lower()
        if not subscription_id:
            continue
        chain = subscription.get("managementGroupAncestors") or []
        # Resource Graph orders immediate parent -> tenant root. Tenant root is
        # always the last ancestor and is intentionally excluded from levels.
        top_down = list(reversed(chain[:-1])) if chain else []
        if output_mode == _MG_SELECTED_LEVEL:
            if len(top_down) < level:
                missing_reasons[subscription_id] = "hierarchy_too_shallow"
                continue
            value = str(top_down[level - 1].get("name") or "").strip()
            if value:
                mapping[subscription_id] = value
            else:
                missing_reasons[subscription_id] = "no_source_value"
            continue

        if not top_down:
            no_path.append(subscription_id)
            continue
        names = [str(node.get("name") or "").strip() for node in top_down[:level]]
        names = [name for name in names if name]
        if names:
            mapping[subscription_id] = " > ".join(names)
        else:
            no_path.append(subscription_id)

    if no_path:
        preview = ", ".join(no_path[:5])
        suffix = "..." if len(no_path) > 5 else ""
        raise ValueError(
            f"Management-group path is unavailable for {len(no_path)} subscription(s): {preview}{suffix}"
        )
    return mapping, missing_reasons


def build_subscription_tag_mapping(
    inventory: Iterable[dict[str, Any]],
    tag_key: str,
) -> tuple[dict[str, str], dict[str, str]]:
    """Resolve one subscription tag key using case-insensitive key matching."""
    requested_key = str(tag_key or "").strip()
    if not 1 <= len(requested_key) <= 512:
        raise ValueError("Subscription tagKey must contain between 1 and 512 characters.")
    requested_casefold = requested_key.casefold()
    mapping: dict[str, str] = {}
    missing_reasons: dict[str, str] = {}
    for subscription in inventory:
        subscription_id = str(subscription.get("id") or "").strip().lower()
        if not subscription_id:
            continue
        tags = subscription.get("tags") if isinstance(subscription.get("tags"), dict) else {}
        value = next(
            (tag_value for key, tag_value in tags.items() if str(key).casefold() == requested_casefold),
            None,
        )
        normalized_value = str(value or "").strip()
        if normalized_value:
            mapping[subscription_id] = normalized_value
        else:
            missing_reasons[subscription_id] = "tag_not_present"
    return mapping, missing_reasons


def rank_subscription_tag_keys(
    inventory: Iterable[dict[str, Any]],
    subscription_ids: Iterable[str],
    limit: int = 15,
) -> list[dict[str, Any]]:
    """Rank tag keys by subscription presence within the supplied scope."""
    allowed = _normalized_ids(subscription_ids)
    counts: dict[str, dict[str, Any]] = {}
    for subscription in inventory:
        subscription_id = str(subscription.get("id") or "").strip().lower()
        if subscription_id not in allowed:
            continue
        tags = subscription.get("tags") if isinstance(subscription.get("tags"), dict) else {}
        seen: set[str] = set()
        for key in tags:
            display_key = str(key).strip()
            normalized_key = display_key.casefold()
            if not display_key or normalized_key in seen:
                continue
            seen.add(normalized_key)
            item = counts.setdefault(normalized_key, {"key": display_key, "subscriptionCount": 0})
            item["subscriptionCount"] += 1
    ranked = sorted(
        counts.values(),
        key=lambda item: (-item["subscriptionCount"], item["key"].casefold()),
    )
    return ranked[:max(limit, 0)]


def build_manual_template(
    inventory: Iterable[dict[str, Any]],
    in_scope_subscription_ids: Iterable[str],
    output_format: str,
    saved_mapping: dict[str, str] | None = None,
) -> tuple[bytes, str, str]:
    """Build a CSV or XLSX template for accessible subscriptions."""
    normalized_format = str(output_format or "").strip().lower()
    if normalized_format not in {"csv", "xlsx"}:
        raise ValueError("Template format must be csv or xlsx.")
    in_scope = _normalized_ids(in_scope_subscription_ids)
    mapping = saved_mapping or {}
    rows = [
        [
            str(item.get("id") or "").strip(),
            str(item.get("name") or "").strip(),
            "Configured scope" if str(item.get("id") or "").strip().lower() in in_scope else "Outside configured scope",
            mapping.get(str(item.get("id") or "").strip().lower(), ""),
        ]
        for item in inventory
        if str(item.get("id") or "").strip()
    ]
    headers = ["SubscriptionId", "SubscriptionName", "Scope", "Value"]
    if normalized_format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(headers)
        writer.writerows([[_csv_safe(value) for value in row] for row in rows])
        return stream.getvalue().encode("utf-8-sig"), "text/csv", "business-context-template.csv"

    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Mapping"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Enter one Value for each subscription to map. Leave Value blank for Not mapped."])
    output = io.BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "business-context-template.xlsx",
    )


def _csv_safe(value: Any) -> Any:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


class SubscriptionContextService:
    def __init__(
        self,
        store: SubscriptionContextStore = context_store,
        max_fields: Any = settings.subscription_context_max_fields,
        refresh_lease_seconds: Any = settings.subscription_context_refresh_lease_seconds,
        inventory_loader: Any = None,
        now: Any = None,
        max_stale_minutes: Any = settings.subscription_context_max_stale_minutes,
    ) -> None:
        self._store = store
        self._max_fields = max_fields
        self._refresh_lease_seconds = refresh_lease_seconds
        self._inventory_loader = inventory_loader or self._load_backend_inventory
        self._now = now or (lambda: datetime.now(timezone.utc))
        self._max_stale_minutes = max_stale_minutes

    def list_fields(self) -> list[dict[str, Any]]:
        return [self._public_definition(item) for item in self._store.list_definitions()]

    def backend_inventory_and_scope(self) -> tuple[list[dict], tuple[str, ...]]:
        backend = for_backend_operation()
        inventory = get_subscription_inventory(backend.get_token())
        app_scope = resolve_app_scope(backend)
        return inventory, app_scope.subscription_ids

    def get_field(self, field_key: str) -> dict[str, Any] | None:
        definition = self._store.get_definition(field_key)
        return self._public_definition(definition) if definition else None

    def get_manual_mapping(self, field_key: str) -> tuple[dict[str, Any], dict[str, str]] | None:
        definition = self._store.get_definition(field_key)
        if definition is None:
            return None
        if definition["valueSource"].get("type") != "manual_file":
            raise ValueError("Context field does not use the manual_file Value source.")
        return self._public_definition(definition), self._store.get_source_mapping(definition)

    def save_manual_field(
        self,
        field: ContextField,
        mapping: dict[str, str],
        updated_by: UserAuditMetadata,
    ) -> dict[str, Any]:
        self._validate_save(field)
        lease = self._publication_lease()
        try:
            saved = self._store.save_manual_definition(
                self._definition_payload(field),
                mapping,
                updated_by,
            )
        finally:
            lease.release()
        return self._public_definition(saved)

    def preview_management_group_field(
        self,
        field: ContextField,
        inventory: Iterable[dict[str, Any]],
    ) -> tuple[dict[str, str], dict[str, str]]:
        if field.value_source.get("type") != _VALUE_SOURCE_MANAGEMENT_GROUP:
            raise ValueError("Context field does not use the management_group_level Value source.")
        config = field.value_source["config"]
        return build_management_group_mapping(inventory, config["level"], config["outputMode"])

    def preview_subscription_tag_field(
        self,
        field: ContextField,
        inventory: Iterable[dict[str, Any]],
    ) -> tuple[dict[str, str], dict[str, str]]:
        if field.value_source.get("type") != _VALUE_SOURCE_SUBSCRIPTION_TAG:
            raise ValueError("Context field does not use the subscription_tag Value source.")
        return build_subscription_tag_mapping(inventory, field.value_source["config"]["tagKey"])

    def save_management_group_field(
        self,
        field: ContextField,
        inventory: Iterable[dict[str, Any]],
        updated_by: UserAuditMetadata,
    ) -> tuple[dict[str, Any], dict[str, str], dict[str, str]]:
        self._validate_save(field)
        mapping, missing_reasons = self.preview_management_group_field(field, inventory)
        resolved_at = self._now()
        lease = self._publication_lease()
        try:
            saved = self._store.save_resolved_definition(
                self._definition_payload(field),
                mapping,
                updated_by,
                updated_at=resolved_at,
            )
        finally:
            lease.release()
        return self._public_definition(saved), mapping, missing_reasons

    def save_subscription_tag_field(
        self,
        field: ContextField,
        inventory: Iterable[dict[str, Any]],
        updated_by: UserAuditMetadata,
    ) -> tuple[dict[str, Any], dict[str, str], dict[str, str]]:
        self._validate_save(field)
        mapping, missing_reasons = self.preview_subscription_tag_field(field, inventory)
        resolved_at = self._now()
        lease = self._publication_lease()
        try:
            saved = self._store.save_resolved_definition(
                self._definition_payload(field),
                mapping,
                updated_by,
                updated_at=resolved_at,
            )
        finally:
            lease.release()
        return self._public_definition(saved), mapping, missing_reasons

    def _validate_save(self, field: ContextField) -> None:
        definitions = self._store.list_definitions()
        existing = next((item for item in definitions if item["key"] == field.key), None)
        if existing is None and len(definitions) >= self._max_fields():
            raise ValueError(f"At most {self._max_fields()} context fields can be configured.")
        duplicate_name = next(
            (
                item for item in definitions
                if item["key"] != field.key and item["name"].casefold() == field.name.casefold()
            ),
            None,
        )
        if duplicate_name:
            raise ValueError(f"Context field name {field.name!r} is already in use.")

    def delete_field(self, field_key: str) -> bool:
        lease = self._publication_lease()
        try:
            return self._store.delete_definition(field_key)
        finally:
            lease.release()

    def _publication_lease(self):
        lease = self._store.try_acquire_operation_lease(self._refresh_lease_seconds())
        if lease is None:
            raise ContextFieldConflictError("Another Business Context operation is already running.")
        return lease

    def resolve(self, subscription_ids, include_disabled: bool = False) -> dict[str, dict[str, str | None]]:
        return self._store.resolve(subscription_ids, include_disabled)

    def has_enabled_dynamic_fields(self) -> bool:
        return any(
            definition.get("enabled")
            and definition.get("valueSource", {}).get("type")
            in {_VALUE_SOURCE_MANAGEMENT_GROUP, _VALUE_SOURCE_SUBSCRIPTION_TAG}
            for definition in self._store.list_definitions()
        )

    def refresh(self, request: dict) -> dict:
        refresh_run_id = str(request.get("refreshRunId") or "").strip()
        if not refresh_run_id:
            return {"status": "error", "error": "refreshRunId is required."}
        lease = self._store.try_acquire_operation_lease(self._refresh_lease_seconds())
        if lease is None:
            # A synchronous save/delete holds the store lease; let the orchestrator retry.
            return {"status": "transient", "error": "Another Business Context operation is in progress."}
        try:
            definitions = self._store.list_definitions()
            scheduled = bool(request.get("scheduled"))
            selected = [
                definition for definition in definitions
                if definition.get("enabled")
                and (
                    not scheduled
                    or definition.get("valueSource", {}).get("type")
                    in {_VALUE_SOURCE_MANAGEMENT_GROUP, _VALUE_SOURCE_SUBSCRIPTION_TAG}
                )
            ]
            if not selected:
                return {"status": "no_work", "fieldCount": 0}

            needs_inventory = any(
                definition["valueSource"]["type"] != _VALUE_SOURCE_MANUAL
                for definition in selected
            )
            inventory = None
            inventory_error = None
            if needs_inventory:
                try:
                    inventory = self._inventory_loader()
                except Exception:  # noqa: BLE001 - recorded as an independent field outcome
                    inventory_error = "Subscription inventory could not be loaded."
            if inventory_error and not request.get("finalAttempt"):
                return {"status": "transient", "error": inventory_error}

            updates: dict[str, dict[str, str]] = {}
            outcomes = []
            blocking_failure = False
            for definition in selected:
                try:
                    source_type = definition["valueSource"]["type"]
                    if source_type == _VALUE_SOURCE_MANUAL:
                        mapping = self._store.get_source_mapping(definition)
                    elif inventory_error:
                        raise RuntimeError(inventory_error)
                    elif source_type == _VALUE_SOURCE_MANAGEMENT_GROUP:
                        config = definition["valueSource"]["config"]
                        mapping, _ = build_management_group_mapping(
                            inventory, config["level"], config["outputMode"]
                        )
                    else:
                        mapping, _ = build_subscription_tag_mapping(
                            inventory, definition["valueSource"]["config"]["tagKey"]
                        )
                    updates[definition["key"]] = mapping
                    outcomes.append({"fieldKey": definition["key"], "status": "ok"})
                except Exception:  # noqa: BLE001 - one field must not abort later fields
                    has_previous = bool(definition.get("sourceBlob"))
                    blocking_failure = blocking_failure or not has_previous
                    outcomes.append({
                        "fieldKey": definition["key"],
                        "status": "stale" if has_previous else "failed",
                        "error": "Source resolution failed.",
                    })

            if blocking_failure:
                return {
                    "status": "failed",
                    "published": False,
                    "fieldCount": len(selected),
                    "outcomes": outcomes,
                }
            publication = self._store.publish_refresh_batch(
                refresh_run_id,
                updates,
                {definition["key"]: definition["revision"] for definition in selected},
                self._now(),
            )
            failures = sum(outcome["status"] != "ok" for outcome in outcomes)
            return {
                "status": "partial" if failures else "completed",
                "published": True,
                "fieldCount": len(selected),
                "failed": failures,
                "outcomes": outcomes,
                **publication,
            }
        finally:
            lease.release()

    def status(self) -> dict:
        definitions = self._store.list_definitions()
        snapshot = self._store.open_snapshot()
        now = self._now()
        fields = (snapshot or {}).get("fields") or {}
        stale = 0
        for metadata in fields.values():
            if not metadata.get("enabled") or metadata.get("valueSourceType") == _VALUE_SOURCE_MANUAL:
                continue
            resolved_at = metadata.get("resolvedAt")
            if resolved_at and now - datetime.fromisoformat(
                str(resolved_at).replace("Z", "+00:00")
            ) > timedelta(minutes=self._max_stale_minutes()):
                stale += 1
        return {
            "hasData": snapshot is not None,
            "generation": snapshot.get("generation") if snapshot else None,
            "publishedAt": snapshot.get("publishedAt") if snapshot else None,
            "configuredFieldCount": len(definitions),
            "enabledFieldCount": sum(bool(item.get("enabled")) for item in definitions),
            "dynamicFieldCount": sum(
                item.get("valueSource", {}).get("type") != _VALUE_SOURCE_MANUAL
                for item in definitions if item.get("enabled")
            ),
            "staleFieldCount": stale,
            "stale": stale > 0,
        }

    def flush(self, _request: dict) -> dict:
        lease = self._store.try_acquire_operation_lease(self._refresh_lease_seconds())
        if lease is None:
            # A synchronous save/delete holds the store lease; let the orchestrator retry.
            return {"status": "transient", "error": "Another Business Context operation is in progress."}
        try:
            return {"status": "completed", "deleted": self._store.flush_generated_data()}
        finally:
            lease.release()

    @staticmethod
    def _definition_payload(field: ContextField) -> dict[str, Any]:
        return {
            "key": field.key,
            "name": field.name,
            "enabled": field.enabled,
            "valueSource": field.value_source,
            "revision": field.revision,
        }

    @staticmethod
    def _load_backend_inventory() -> list[dict[str, Any]]:
        backend = for_backend_operation()
        return get_subscription_inventory(backend.get_token())

    @staticmethod
    def _public_definition(definition: dict[str, Any]) -> dict[str, Any]:
        return {
            key: value.isoformat() if hasattr(value, "isoformat") else value
            for key, value in definition.items()
            if key not in {"etag", "sourceBlob"}
        }


service = SubscriptionContextService()


def enrich_rows_with_business_context(
    rows: Iterable[dict[str, Any]],
    context_service: SubscriptionContextService = service,
) -> dict[str, str]:
    """Attach enabled context values to rows and return key/display-name metadata."""
    materialized_rows = list(rows)
    fields = {
        field["key"]: field["name"]
        for field in context_service.list_fields()
        if field.get("enabled")
    }
    subscription_ids = sorted(_normalized_ids(
        row.get("subscriptionId") for row in materialized_rows if row.get("subscriptionId")
    ))
    values = context_service.resolve(subscription_ids) if subscription_ids else {}
    for row in materialized_rows:
        subscription_id = str(row.get("subscriptionId") or "").strip().lower()
        subscription_values = values.get(subscription_id) or {}
        row["businessContext"] = {
            field_key: subscription_values.get(field_key)
            for field_key in fields
        }
    return fields
