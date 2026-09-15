from __future__ import annotations

import io
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import MagicMock

from openpyxl import Workbook

from services.subscription_context import (
    ContextField,
    SubscriptionContextService,
    build_manual_template,
    build_management_group_mapping,
    build_preview,
    build_subscription_tag_mapping,
    parse_context_field,
    parse_manual_mapping,
    rank_subscription_tag_keys,
    enrich_rows_with_business_context,
)

SUB_A = "11111111-1111-1111-1111-111111111111"
SUB_B = "22222222-2222-2222-2222-222222222222"
SUB_C = "33333333-3333-3333-3333-333333333333"


class SubscriptionContextTests(unittest.TestCase):
    def test_parse_context_field_normalizes_key(self):
        field = parse_context_field({
            "key": " Business_Unit ",
            "name": " Business Unit ",
            "enabled": True,
            "valueSource": {"type": "manual_file", "version": 1},
        })

        self.assertEqual("business_unit", field.key)
        self.assertEqual("Business Unit", field.name)

    def test_parse_context_field_rejects_reserved_key(self):
        with self.assertRaisesRegex(ValueError, "reserved"):
            parse_context_field({
                "key": "location",
                "name": "Location",
                "enabled": True,
                "valueSource": {"type": "manual_file", "version": 1},
            })

    def test_parse_management_group_field(self):
        field = parse_context_field({
            "key": "organization",
            "name": "Organization",
            "enabled": True,
            "valueSource": {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 3, "outputMode": "hierarchy_path"},
            },
        })

        self.assertEqual(3, field.value_source["config"]["level"])

    def test_parse_management_group_field_rejects_level_outside_fixed_range(self):
        with self.assertRaisesRegex(ValueError, "1 through 6"):
            parse_context_field({
                "key": "organization",
                "name": "Organization",
                "enabled": True,
                "valueSource": {
                    "type": "management_group_level",
                    "version": 1,
                    "config": {"level": 7, "outputMode": "selected_level"},
                },
            })

    def test_parse_subscription_tag_field(self):
        field = parse_context_field({
            "key": "environment",
            "name": "Environment",
            "enabled": True,
            "valueSource": {
                "type": "subscription_tag",
                "version": 1,
                "config": {"tagKey": " Environment "},
            },
        })

        self.assertEqual(" Environment ", field.value_source["config"]["tagKey"])

    def test_subscription_tag_mapping_matches_key_case_insensitively(self):
        mapping, missing = build_subscription_tag_mapping(
            [
                {"id": SUB_A, "tags": {"Environment": "PRD"}},
                {"id": SUB_B, "tags": {"ENVIRONMENT": ""}},
                {"id": SUB_C, "tags": {}},
            ],
            "environment",
        )

        self.assertEqual({SUB_A: "PRD"}, mapping)
        self.assertEqual("tag_not_present", missing[SUB_B])
        self.assertEqual("tag_not_present", missing[SUB_C])

    def test_rank_subscription_tag_keys_uses_only_supplied_scope(self):
        ranked = rank_subscription_tag_keys(
            [
                {"id": SUB_A, "tags": {"Environment": "PRD", "CostCenter": "1"}},
                {"id": SUB_B, "tags": {"environment": "DEV"}},
                {"id": SUB_C, "tags": {"OutsideOnly": "x"}},
            ],
            [SUB_A, SUB_B],
        )

        self.assertEqual(
            [
                {"key": "Environment", "subscriptionCount": 2},
                {"key": "CostCenter", "subscriptionCount": 1},
            ],
            ranked,
        )

    def test_parse_csv_rejects_duplicate_subscription(self):
        content = (
            "SubscriptionId,Value\n"
            f"{SUB_A},HR\n"
            f"{SUB_A},IT\n"
        ).encode()

        with self.assertRaisesRegex(ValueError, "Duplicate SubscriptionId"):
            parse_manual_mapping(content, "mapping.csv")

    def test_parse_csv_keeps_unknown_and_omits_blank_value(self):
        content = (
            "SubscriptionId,Value\n"
            f"{SUB_A},Unknown\n"
            f"{SUB_B},\n"
        ).encode()

        mapping = parse_manual_mapping(content, "mapping.csv")

        self.assertEqual({SUB_A: "Unknown"}, mapping)

    def test_parse_xlsx_rejects_formulas(self):
        content = self._workbook_bytes([
            ["SubscriptionId", "Value"],
            [SUB_A, "=UPPER(\"hr\")"],
        ])

        with self.assertRaisesRegex(ValueError, "must not contain formulas"):
            parse_manual_mapping(content, "mapping.xlsx")

    def test_parse_xlsx_mapping(self):
        content = self._workbook_bytes([
            ["SubscriptionId", "Value"],
            [SUB_A.upper(), "HR"],
            [SUB_B, None],
        ])

        mapping = parse_manual_mapping(content, "mapping.xlsx")

        self.assertEqual({SUB_A: "HR"}, mapping)

    def test_build_preview_separates_scope_and_null_values(self):
        preview = build_preview(
            {SUB_A: "Unknown", SUB_C: "IT"},
            [SUB_A, SUB_B, SUB_C],
            [SUB_A, SUB_B],
        )

        self.assertEqual("Unknown", preview["values"][SUB_A])
        self.assertIsNone(preview["values"][SUB_B])
        self.assertEqual(1, preview["inScope"]["mappedSubscriptions"])
        self.assertEqual(1, preview["inScope"]["unmappedSubscriptions"])
        self.assertEqual(50.0, preview["inScope"]["coveragePercent"])
        self.assertEqual(1, preview["outsideScope"]["mappedSubscriptions"])

    def test_management_group_selected_level_uses_root_down_order(self):
        mapping, missing = build_management_group_mapping(
            [
                self._subscription(SUB_A, "Team", "Division", "Organization", "Tenant Root"),
                self._subscription(SUB_B, "Organization", "Tenant Root"),
            ],
            level=2,
            output_mode="selected_level",
        )

        self.assertEqual("Division", mapping[SUB_A])
        self.assertNotIn(SUB_B, mapping)
        self.assertEqual("hierarchy_too_shallow", missing[SUB_B])

    def test_management_group_path_truncates_to_available_depth(self):
        mapping, missing = build_management_group_mapping(
            [
                self._subscription(SUB_A, "Team", "Division", "Organization", "Tenant Root"),
                self._subscription(SUB_B, "Organization", "Tenant Root"),
            ],
            level=3,
            output_mode="hierarchy_path",
        )

        self.assertEqual("Organization > Division > Team", mapping[SUB_A])
        self.assertEqual("Organization", mapping[SUB_B])
        self.assertEqual({}, missing)

    def test_management_group_path_fails_when_non_root_path_is_absent(self):
        with self.assertRaisesRegex(ValueError, "path is unavailable"):
            build_management_group_mapping(
                [self._subscription(SUB_A, "Tenant Root")],
                level=1,
                output_mode="hierarchy_path",
            )

    def test_build_creation_csv_template(self):
        content, mimetype, filename = build_manual_template(
            [{"id": SUB_A, "name": "Subscription A"}, {"id": SUB_B, "name": "Subscription B"}],
            [SUB_A],
            "csv",
        )

        text = content.decode("utf-8-sig")
        self.assertEqual("text/csv", mimetype)
        self.assertEqual("business-context-template.csv", filename)
        self.assertIn(f"{SUB_A},Subscription A,Configured scope,", text)
        self.assertIn(f"{SUB_B},Subscription B,Outside configured scope,", text)

    def test_build_prefilled_xlsx_template(self):
        content, mimetype, filename = build_manual_template(
            [{"id": SUB_A, "name": "Subscription A"}],
            [SUB_A],
            "xlsx",
            {SUB_A: "HR"},
        )

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True)
        rows = list(workbook["Mapping"].iter_rows(values_only=True))
        self.assertEqual("HR", rows[1][3])
        self.assertEqual(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            mimetype,
        )
        self.assertEqual("business-context-template.xlsx", filename)

    def test_service_enforces_maximum_field_count(self):
        store = MagicMock()
        store.list_definitions.return_value = [
            {"key": "existing", "name": "Existing"},
        ]
        service = SubscriptionContextService(store, lambda: 1)

        with self.assertRaisesRegex(ValueError, "At most 1"):
            service.save_manual_field(
                ContextField("new_field", "New field", True, {"type": "manual_file", "version": 1}),
                {},
                MagicMock(),
            )

    def test_service_rejects_duplicate_name_case_insensitively(self):
        store = MagicMock()
        store.list_definitions.return_value = [
            {"key": "existing", "name": "Business Unit"},
        ]
        service = SubscriptionContextService(store)

        with self.assertRaisesRegex(ValueError, "already in use"):
            service.save_manual_field(
                ContextField("other", "business unit", True, {"type": "manual_file", "version": 1}),
                {},
                MagicMock(),
            )

    def test_enrich_rows_with_business_context_uses_enabled_fields_and_normalized_ids(self):
        context_service = MagicMock()
        context_service.list_fields.return_value = [
            {"key": "environment", "name": "Environment", "enabled": True},
            {"key": "disabled", "name": "Disabled", "enabled": False},
        ]
        context_service.resolve.return_value = {SUB_A: {"environment": None}}

        rows = [{"subscriptionId": SUB_A.upper()}, {"subscriptionId": SUB_A}]
        metadata = enrich_rows_with_business_context(rows, context_service)

        self.assertEqual({"environment": "Environment"}, metadata)
        self.assertEqual({"environment": None}, rows[0]["businessContext"])
        self.assertEqual({"environment": None}, rows[1]["businessContext"])
        context_service.resolve.assert_called_once_with([SUB_A])

    def test_service_saves_management_group_field_without_read_expiry(self):
        store = MagicMock()
        store.list_definitions.return_value = []
        now = datetime(2026, 8, 24, 12, 0, tzinfo=timezone.utc)
        store.save_resolved_definition.side_effect = lambda definition, mapping, audit, **kwargs: {
            **definition,
            "revision": 1,
            "sourceBlob": "source",
            "updatedAt": kwargs["updated_at"],
            "updatedBy": {},
            "expiresAt": kwargs.get("expires_at"),
        }
        service = SubscriptionContextService(store, now=lambda: now)
        field = ContextField(
            "organization",
            "Organization",
            True,
            {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 2, "outputMode": "selected_level"},
            },
        )

        saved, mapping, missing = service.save_management_group_field(
            field,
            [self._subscription(SUB_A, "Team", "Division", "Organization", "Tenant Root")],
            MagicMock(),
        )

        self.assertEqual("Division", mapping[SUB_A])
        self.assertEqual({}, missing)
        self.assertIsNone(saved["expiresAt"])

    def test_service_saves_subscription_tag_field_without_read_expiry(self):
        store = MagicMock()
        store.list_definitions.return_value = []
        now = datetime(2026, 8, 24, 12, 0, tzinfo=timezone.utc)
        store.save_resolved_definition.side_effect = lambda definition, mapping, audit, **kwargs: {
            **definition,
            "revision": 1,
            "sourceBlob": "source",
            "updatedAt": kwargs["updated_at"],
            "updatedBy": {},
            "expiresAt": kwargs.get("expires_at"),
        }
        service = SubscriptionContextService(store, now=lambda: now)
        field = ContextField(
            "environment",
            "Environment",
            True,
            {
                "type": "subscription_tag",
                "version": 1,
                "config": {"tagKey": "Environment"},
            },
        )

        saved, mapping, missing = service.save_subscription_tag_field(
            field,
            [{"id": SUB_A, "tags": {"environment": "PRD"}}, {"id": SUB_B, "tags": {}}],
            MagicMock(),
        )

        self.assertEqual({SUB_A: "PRD"}, mapping)
        self.assertEqual("tag_not_present", missing[SUB_B])
        self.assertIsNone(saved["expiresAt"])

    def legacy_resolve_refreshes_expired_subscription_tag_mapping(self):
        store = MagicMock()
        now = datetime(2026, 8, 24, 12, 5, tzinfo=timezone.utc)
        definition = {
            "key": "environment",
            "name": "Environment",
            "enabled": True,
            "revision": 2,
            "valueSource": {
                "type": "subscription_tag",
                "version": 1,
                "config": {"tagKey": "Environment"},
            },
            "expiresAt": now - timedelta(seconds=1),
            "resolvedAt": now - timedelta(minutes=5),
        }
        store.list_definitions.return_value = [definition]
        store.get_definition.return_value = definition
        lease = MagicMock()
        store.try_acquire_refresh_lease.return_value = lease
        store.resolve.return_value = {SUB_A: {"environment": "PRD"}}
        service = SubscriptionContextService(
            store,
            inventory_loader=lambda: [{"id": SUB_A, "tags": {"environment": "PRD"}}],
            now=lambda: now,
        )

        result = service.resolve([SUB_A])

        self.assertEqual("PRD", result[SUB_A]["environment"])
        store.refresh_mapping.assert_called_once_with(
            "environment",
            {SUB_A: "PRD"},
            now,
            now + timedelta(minutes=5),
            expected_revision=2,
        )

    def legacy_resolve_refreshes_expired_management_group_mapping(self):
        store = MagicMock()
        now = datetime(2026, 8, 24, 12, 5, tzinfo=timezone.utc)
        store.list_definitions.return_value = [{
            "key": "organization",
            "name": "Organization",
            "enabled": True,
            "revision": 3,
            "valueSource": {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 1, "outputMode": "selected_level"},
            },
            "expiresAt": now - timedelta(seconds=1),
            "resolvedAt": now - timedelta(minutes=5),
        }]
        store.get_definition.return_value = store.list_definitions.return_value[0]
        lease = MagicMock()
        store.try_acquire_refresh_lease.return_value = lease
        store.resolve.return_value = {SUB_A: {"organization": "Organization"}}
        inventory_loader = MagicMock(return_value=[
            self._subscription(SUB_A, "Team", "Organization", "Tenant Root"),
        ])
        service = SubscriptionContextService(store, inventory_loader=inventory_loader, now=lambda: now)

        result = service.resolve([SUB_A])

        self.assertEqual("Organization", result[SUB_A]["organization"])
        inventory_loader.assert_called_once_with()
        store.refresh_mapping.assert_called_once_with(
            "organization",
            {SUB_A: "Organization"},
            now,
            now + timedelta(minutes=5),
            expected_revision=3,
        )
        lease.renew.assert_called_once_with()
        lease.release.assert_called_once_with()

    def legacy_resolve_contender_waits_for_published_refresh_without_loading_inventory(self):
        store = MagicMock()
        now = datetime(2026, 8, 24, 12, 5, tzinfo=timezone.utc)
        stale = {
            "key": "organization",
            "name": "Organization",
            "enabled": True,
            "revision": 3,
            "valueSource": {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 1, "outputMode": "selected_level"},
            },
            "resolvedAt": now - timedelta(minutes=5),
            "expiresAt": now - timedelta(seconds=1),
        }
        fresh = {**stale, "resolvedAt": now, "expiresAt": now + timedelta(minutes=5)}
        store.list_definitions.return_value = [stale]
        store.try_acquire_refresh_lease.return_value = None
        store.get_definition.return_value = fresh
        store.snapshot_field_is_fresh.return_value = True
        store.resolve.return_value = {SUB_A: {"organization": "Organization"}}
        inventory_loader = MagicMock()
        service = SubscriptionContextService(store, inventory_loader=inventory_loader, now=lambda: now)

        result = service.resolve([SUB_A])

        self.assertEqual("Organization", result[SUB_A]["organization"])
        inventory_loader.assert_not_called()
        store.refresh_mapping.assert_not_called()
        store.snapshot_field_is_fresh.assert_called_once_with(
            "organization",
            3,
            stale["resolvedAt"],
        )

    def legacy_refresh_owner_releases_lease_when_inventory_fails(self):
        store = MagicMock()
        now = datetime(2026, 8, 24, 12, 5, tzinfo=timezone.utc)
        stale = {
            "key": "organization",
            "enabled": True,
            "revision": 3,
            "valueSource": {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 1, "outputMode": "selected_level"},
            },
            "resolvedAt": now - timedelta(minutes=5),
            "expiresAt": now - timedelta(seconds=1),
        }
        store.list_definitions.return_value = [stale]
        store.get_definition.return_value = stale
        lease = MagicMock()
        store.try_acquire_refresh_lease.return_value = lease
        inventory_loader = MagicMock(side_effect=RuntimeError("Resource Graph unavailable"))
        service = SubscriptionContextService(store, inventory_loader=inventory_loader, now=lambda: now)

        with self.assertRaisesRegex(RuntimeError, "Resource Graph unavailable"):
            service.resolve([SUB_A])

        lease.release.assert_called_once_with()
        lease.renew.assert_not_called()
        store.refresh_mapping.assert_not_called()

    def legacy_refresh_contender_times_out_without_loading_inventory(self):
        store = MagicMock()
        now = datetime(2026, 8, 24, 12, 5, tzinfo=timezone.utc)
        stale = {
            "key": "organization",
            "enabled": True,
            "revision": 3,
            "valueSource": {
                "type": "management_group_level",
                "version": 1,
                "config": {"level": 1, "outputMode": "selected_level"},
            },
            "resolvedAt": now - timedelta(minutes=5),
            "expiresAt": now - timedelta(seconds=1),
        }
        store.list_definitions.return_value = [stale]
        store.get_definition.return_value = stale
        store.try_acquire_refresh_lease.return_value = None
        inventory_loader = MagicMock()
        monotonic = MagicMock(side_effect=[0.0, 2.0])
        sleeper = MagicMock()
        service = SubscriptionContextService(
            store,
            refresh_wait_seconds=lambda: 1,
            inventory_loader=inventory_loader,
            now=lambda: now,
            monotonic=monotonic,
            sleeper=sleeper,
        )

        with self.assertRaisesRegex(TimeoutError, "Timed out waiting"):
            service.resolve([SUB_A])

        inventory_loader.assert_not_called()
        sleeper.assert_not_called()

    @staticmethod
    def _workbook_bytes(rows) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        for row in rows:
            worksheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @staticmethod
    def _subscription(subscription_id: str, *ancestors_immediate_to_root: str):
        return {
            "id": subscription_id,
            "managementGroupAncestors": [
                {"id": name.lower().replace(" ", "_"), "name": name}
                for name in ancestors_immediate_to_root
            ],
        }


if __name__ == "__main__":
    unittest.main()